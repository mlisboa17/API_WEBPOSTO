"""Cruzamento planilha-alvo × NF-e importadas × webPosto (GET) × modelos fiscais."""

from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.operational.dfe import store
from src.operational.dfe.packaging import resolve_packaging
from src.operational.dfe.text_fix import fix_mojibake, fix_tree
from src.operational.price_update.value_normalizer import (
    normalize_ean,
    normalize_name,
    parse_price_br,
)

DEFAULT_COMPANY = 118508
DEFAULT_UF = "PE"
DEFAULT_REGIME = "LUCRO_PRESUMIDO"
DEFAULT_COST_CENTER = 24886
DEFAULT_PRICE_TABLE = "A"

STATUS_PT = {
    "ALREADY_REGISTERED": "Já cadastrado",
    "INACTIVE_EXISTING": "Cadastrado (inativo)",
    "POSSIBLE_DUPLICATE": "Possível duplicidade",
    "NEW_PRODUCT": "Novo produto",
    "AMBIGUOUS": "Ambíguo no webPosto",
    "INVALID_EAN": "EAN inválido",
    "XML_EXACT_MATCH": "Match exato na NF-e",
    "XML_PACKAGING_MATCH": "Match por embalagem",
    "XML_DESCRIPTION_MATCH": "Match por descrição",
    "XML_AMBIGUOUS": "Vários candidatos na NF-e",
    "XML_NOT_FOUND": "Sem evidência na NF-e",
    "READY_TO_CREATE": "Pronto para cadastrar",
    "READY_FOR_REVIEW": "Pronto para revisão",
    "NO_XML_EVIDENCE": "Sem XML",
    "NO_FISCAL_MODEL": "Sem modelo fiscal",
    "PACKAGING_AMBIGUITY": "Conferir embalagem",
    "TAX_CONFLICT": "Conflito tributário",
    "MISSING_PURCHASE_PRICE": "Sem preço de compra",
    "INCOMPLETE_BODY": "Cadastro incompleto",
    "INVALID_ROW": "Linha inválida",
    "BLOCKED": "Bloqueado",
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _norm_header(h: str) -> str:
    import unicodedata

    s = unicodedata.normalize("NFKD", h or "")
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"[\s_\-]+", " ", s.lower()).strip()


def suggest_target_mapping(headers: list[str]) -> dict[str, str | None]:
    def find(*cands: str) -> str | None:
        norms = {_norm_header(h): h for h in headers if h}
        # 1) match exato normalizado
        for c in cands:
            cn = _norm_header(c)
            if cn in norms:
                return norms[cn]
        # 2) contains (mais específico primeiro)
        for c in cands:
            cn = _norm_header(c)
            for hn, h in norms.items():
                if cn and (cn == hn or cn in hn or hn in cn):
                    return h
        return None

    return {
        "ean": find(
            "ean13",
            "ean",
            "gtin",
            "codigo de barras",
            "código de barras",
            "codigo barras",
            "codbarras",
            "cod barras",
            "codigo_barras",
        ),
        "descricao": find(
            "xprod",
            "descricao",
            "descrição",
            "produto",
            "nome",
            "item",
        ),
        "preco": find(
            "preco venda",
            "preço venda",
            "valor de venda",
            "valor venda",
            "preco_venda",
            "preço",
            "preco",
            "price",
            "valor",
        ),
    }


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def list_sheet_names(data: bytes, filename: str) -> list[str]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return ["CSV"]
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        try:
            import xlrd  # type: ignore
        except ImportError as exc:
            raise ValueError(
                "XLS legado requer pacote xlrd — salve como XLSX ou CSV"
            ) from exc
        book = xlrd.open_workbook(file_contents=data)
        return list(book.sheet_names())
    bio = io.BytesIO(data)
    wb = load_workbook(bio, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _decode_csv_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _cell_str(value: Any) -> str:
    """Texto de célula preservando EAN (evita float científico)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        # notação científica → dígitos inteiros se possível
        as_int = int(round(value))
        if abs(value - as_int) < 1e-9:
            return str(as_int)
        return format(value, "f").rstrip("0").rstrip(".")
    return str(value).strip()


def _read_matrix(
    data: bytes,
    filename: str,
    sheet: str | None,
    *,
    header_row: int = 1,
) -> tuple[list[str], list[list[Any]], str]:
    """Lê matriz. header_row é 1-based (linha do cabeçalho na planilha)."""
    name = (filename or "").lower()
    hr = max(1, int(header_row or 1))

    if name.endswith(".csv"):
        text = _decode_csv_text(data)
        try:
            sample = text[:4096]
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";" if text.count(";") >= text.count(",") else ","
        reader = csv.reader(io.StringIO(text), dialect)
        rows = [list(r) for r in reader]
        if not rows or len(rows) < hr:
            return [], [], "CSV"
        headers = [_cell_str(c) for c in rows[hr - 1]]
        body = [list(r) for r in rows[hr:]]
        return headers, body, "CSV"

    if name.endswith(".xls") and not name.endswith(".xlsx"):
        try:
            import xlrd  # type: ignore
        except ImportError as exc:
            raise ValueError(
                "XLS legado requer pacote xlrd — salve como XLSX ou CSV"
            ) from exc
        book = xlrd.open_workbook(file_contents=data)
        sheet_name = sheet or book.sheet_names()[0]
        if sheet_name not in book.sheet_names():
            raise ValueError(f"Aba não encontrada: {sheet_name}")
        ws = book.sheet_by_name(sheet_name)
        if ws.nrows < hr:
            return [], [], sheet_name
        headers = [_cell_str(ws.cell_value(hr - 1, c)) for c in range(ws.ncols)]
        body = [
            [_cell_str(ws.cell_value(r, c)) for c in range(ws.ncols)]
            for r in range(hr, ws.nrows)
        ]
        return headers, body, sheet_name

    bio = io.BytesIO(data)
    wb = load_workbook(bio, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        if not names:
            return [], [], sheet or ""
        sheet_name = sheet or names[0]
        if sheet_name not in names:
            raise ValueError(f"Aba não encontrada: {sheet_name}. Disponíveis: {names}")
        ws = wb[sheet_name]
        rows_all = [list(r) for r in ws.iter_rows(values_only=True)]
        if len(rows_all) < hr:
            return [], [], sheet_name
        headers = [_cell_str(c) for c in rows_all[hr - 1]]
        body = [list(r) for r in rows_all[hr:]]
        return headers, body, sheet_name
    finally:
        wb.close()


def inspect_target_spreadsheet(
    data: bytes,
    filename: str,
    *,
    sheet: str | None = None,
    header_row: int = 1,
) -> dict[str, Any]:
    """Inspeção de abas/cabeçalhos/mapeamento — sem WebPosto, sem SEFAZ."""
    if not data:
        raise ValueError("Arquivo vazio")
    name = filename or "planilha.xlsx"
    lower = name.lower()
    if not (lower.endswith(".xlsx") or lower.endswith(".xls") or lower.endswith(".csv")):
        raise ValueError("Extensão não suportada. Use XLSX, XLS ou CSV.")

    sheets = list_sheet_names(data, name)
    if not sheets:
        raise ValueError("Nenhuma aba encontrada na planilha")

    summaries: list[dict[str, Any]] = []
    for sn in sheets:
        try:
            headers_s, body_s, used = _read_matrix(data, name, sn if sn != "CSV" else None, header_row=header_row)
            non_empty = any(h.strip() for h in headers_s) or any(
                any(_cell_str(c) for c in row) for row in body_s[:5]
            )
            summaries.append(
                {
                    "name": used,
                    "empty": not non_empty,
                    "headerCount": sum(1 for h in headers_s if h.strip()),
                    "dataRowsSample": min(len(body_s), 5),
                }
            )
        except Exception as exc:
            summaries.append({"name": sn, "empty": True, "error": str(exc)[:120]})

    selected = sheet
    if selected and selected not in sheets and selected != "CSV":
        raise ValueError(f"Aba '{selected}' não existe. Disponíveis: {sheets}")
    if not selected:
        for s in summaries:
            if not s.get("empty") and not s.get("error"):
                selected = s["name"]
                break
        if not selected:
            selected = sheets[0]

    headers, _body, used_sheet = _read_matrix(
        data, name, selected if selected != "CSV" else None, header_row=header_row
    )
    # remove cabeçalhos totalmente vazios do fim
    while headers and not headers[-1].strip():
        headers.pop()
    mapping = suggest_target_mapping(headers)
    return {
        "success": True,
        "filename": name,
        "size": len(data),
        "sha256": sha256_bytes(data),
        "sheets": sheets,
        "sheetSummaries": summaries,
        "selectedSheet": used_sheet,
        "headerRow": header_row,
        "headers": headers,
        "suggestedMapping": mapping,
        "mappingReady": bool(mapping.get("ean") and mapping.get("descricao") and mapping.get("preco")),
        "status": "COLUMN_MAPPING_READY"
        if mapping.get("ean") and mapping.get("descricao") and mapping.get("preco")
        else ("SPREADSHEET_HEADERS_VISIBLE" if headers else "SPREADSHEET_SHEETS_VISIBLE"),
        "webpostoWrites": 0,
        "sefazQueries": 0,
        "analysisStatus": "SHEET_INSPECTED",
        "message": (
            f"Planilha lida: {len(sheets)} aba(s), cabeçalho na linha {header_row}, "
            f"{len([h for h in headers if h])} coluna(s)."
        ),
    }


def parse_target_spreadsheet(
    data: bytes,
    filename: str,
    *,
    sheet: str | None = None,
    col_ean: str | None = None,
    col_desc: str | None = None,
    col_price: str | None = None,
    header_row: int = 1,
) -> dict[str, Any]:
    headers, body, used_sheet = _read_matrix(data, filename, sheet, header_row=header_row)
    mapping = suggest_target_mapping(headers)
    ean_h = col_ean or mapping["ean"]
    desc_h = col_desc or mapping["descricao"]
    price_h = col_price or mapping["preco"]
    if not ean_h or not desc_h or not price_h:
        raise ValueError(
            f"Cabeçalhos obrigatórios não encontrados. Detectado: {mapping}. Headers: {headers}"
        )

    idx = {h: i for i, h in enumerate(headers)}

    def cell(row: list[Any], h: str) -> Any:
        i = idx.get(h)
        if i is None or i >= len(row):
            return None
        return row[i]

    lines: list[dict[str, Any]] = []
    seen_ean: dict[str, int] = {}
    for n, row in enumerate(body, start=2):
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        ean, ean_issues = normalize_ean(cell(row, ean_h))
        # preservar zeros: se Excel entregou float, normalize_ean bloqueia .0
        raw_ean = cell(row, ean_h)
        if ean is None and raw_ean is not None and not ean_issues:
            # tentar como texto puro de dígitos
            digits = re.sub(r"\D", "", str(raw_ean))
            if digits:
                ean, ean_issues = normalize_ean(digits)
        desc = fix_mojibake(str(cell(row, desc_h)).strip() if cell(row, desc_h) is not None else "")
        price, price_raw, price_issues = parse_price_br(cell(row, price_h))
        dup_line = seen_ean.get(ean) if ean else None
        if ean:
            seen_ean.setdefault(ean, n)
        lines.append(
            {
                "row": n,
                "ean": ean,
                "eanRaw": str(raw_ean) if raw_ean is not None else None,
                "eanIssues": ean_issues,
                "descricao": desc or None,
                "descricaoNorm": normalize_name(desc) if desc else "",
                "precoVenda": str(price) if price is not None else None,
                "precoRaw": price_raw,
                "precoIssues": price_issues,
                "duplicateInSheetRow": dup_line if dup_line != n else None,
            }
        )

    return {
        "filename": filename,
        "sheet": used_sheet,
        "size": len(data),
        "sha256": sha256_bytes(data),
        "headers": headers,
        "mapping": {"ean": ean_h, "descricao": desc_h, "preco": price_h},
        "lineCount": len(lines),
        "lines": lines,
    }


def _load_nfe_index(company_code: int) -> dict[str, Any]:
    docs = store.list_documents(company_code)
    # Reutilizar apenas NF-e válidas já importadas (não canceladas / não inválidas).
    docs = [
        d
        for d in docs
        if d.get("status") not in {"CANCELLED", "INVALID", "QUARANTINE", "FAILED"}
    ]
    items_flat: list[dict[str, Any]] = []
    by_ean: dict[str, list[dict[str, Any]]] = {}
    by_ean_trib: dict[str, list[dict[str, Any]]] = {}
    by_name: dict[str, list[dict[str, Any]]] = {}

    for doc in docs:
        its = store.load_items(doc["id"])
        for it in its:
            pack = it.get("packaging") or resolve_packaging(it)
            entry = {
                "documentId": doc["id"],
                "accessKeyMasked": doc.get("access_key_masked"),
                "issuedAt": doc.get("issued_at"),
                "issuerName": fix_mojibake(doc.get("issuer_name")),
                "issuerCnpj": doc.get("issuer_cnpj"),
                "nNF": doc.get("nNF"),
                "serie": doc.get("serie"),
                "docStatus": doc.get("status"),
                "item": {
                    **it,
                    "x_prod": fix_mojibake(it.get("x_prod")),
                    "packaging": pack,
                },
            }
            items_flat.append(entry)
            ean = it.get("c_ean")
            ean_t = it.get("c_ean_trib")
            if ean:
                by_ean.setdefault(str(ean), []).append(entry)
            if ean_t:
                by_ean_trib.setdefault(str(ean_t), []).append(entry)
            nn = normalize_name(it.get("x_prod"))
            if nn:
                by_name.setdefault(nn, []).append(entry)

    return {
        "documents": docs,
        "documentCount": len(docs),
        "itemCount": len(items_flat),
        "items": items_flat,
        "byEan": by_ean,
        "byEanTrib": by_ean_trib,
        "byName": by_name,
    }


def _pick_best_nfe(candidates: list[dict[str, Any]]) -> dict[str, Any] | None:
    if not candidates:
        return None
    ranked = sorted(
        candidates,
        key=lambda c: str(c.get("issuedAt") or ""),
        reverse=True,
    )
    return ranked[0]


def _match_nfe(line: dict[str, Any], nfe_idx: dict[str, Any]) -> dict[str, Any]:
    ean = line.get("ean")
    hits: list[dict[str, Any]] = []
    kind = "XML_NOT_FOUND"
    if ean:
        trib = list(nfe_idx["byEanTrib"].get(ean, []))
        com = list(nfe_idx["byEan"].get(ean, []))
        if trib:
            hits = trib
            kind = "XML_EXACT_MATCH"
        elif com:
            hits = com
            # se uCom sugere embalagem e ean == cEAN (caixa)
            kind = "XML_PACKAGING_MATCH"
            for h in com:
                it = h["item"]
                if it.get("c_ean_trib") and it.get("c_ean") and it["c_ean"] != it["c_ean_trib"]:
                    kind = "XML_PACKAGING_MATCH"
                    break
                if it.get("c_ean") == ean and (it.get("c_ean_trib") in (None, ean)):
                    kind = "XML_EXACT_MATCH"
        # relação cruzada: planilha EAN unidade vs caixa na NF
        if not hits:
            for entry in nfe_idx["items"]:
                it = entry["item"]
                if ean in {it.get("c_ean"), it.get("c_ean_trib")}:
                    hits.append(entry)
            if hits:
                kind = "XML_PACKAGING_MATCH"

    if not hits and line.get("descricaoNorm"):
        name_hits = nfe_idx["byName"].get(line["descricaoNorm"], [])
        if len(name_hits) == 1:
            hits = name_hits
            kind = "XML_DESCRIPTION_MATCH"
        elif len(name_hits) > 1:
            return {
                "xmlStatus": "XML_AMBIGUOUS",
                "xmlStatusLabel": STATUS_PT["XML_AMBIGUOUS"],
                "candidates": len(name_hits),
                "chosen": None,
                "history": [],
            }

    # desambiguar por EAN
    if len(hits) > 1 and ean:
        # se todos mesmos produto (mesmo ean pair) ok
        keys = {
            (h["item"].get("c_ean"), h["item"].get("c_ean_trib"), h["item"].get("x_prod"))
            for h in hits
        }
        if len(keys) > 1 and kind != "XML_EXACT_MATCH":
            # ainda pode ser histórico do mesmo item
            eans_set = {(h["item"].get("c_ean"), h["item"].get("c_ean_trib")) for h in hits}
            if len(eans_set) > 1:
                return {
                    "xmlStatus": "XML_AMBIGUOUS",
                    "xmlStatusLabel": STATUS_PT["XML_AMBIGUOUS"],
                    "candidates": len(hits),
                    "chosen": None,
                    "history": [
                        {
                            "documentId": h["documentId"],
                            "nNF": h.get("nNF"),
                            "issuedAt": h.get("issuedAt"),
                            "xProd": h["item"].get("x_prod"),
                        }
                        for h in hits[:8]
                    ],
                }

    if not hits:
        return {
            "xmlStatus": "XML_NOT_FOUND",
            "xmlStatusLabel": STATUS_PT["XML_NOT_FOUND"],
            "candidates": 0,
            "chosen": None,
            "history": [],
        }

    chosen = _pick_best_nfe(hits)
    it = chosen["item"]
    pack = it.get("packaging") or resolve_packaging(it)
    v_un = it.get("v_un_com")
    factor = pack.get("conversion_factor")
    unit_cost = pack.get("purchase_unit_cost")
    factor_f = float(factor) if factor is not None else None
    content_label = None
    if factor_f is not None and factor_f > 1:
        content_label = (
            f"{int(factor_f) if factor_f.is_integer() else factor_f} UNIDADES"
        )
    elif factor_f == 1:
        content_label = "1 UNIDADE"
    explain = {
        "comprado": f"1 {it.get('u_com') or 'UN'} por R$ {Decimal(str(v_un)):.2f}" if v_un is not None else None,
        "conteudo": content_label,
        "custoUnitario": f"R$ {Decimal(str(unit_cost)):.2f}" if unit_cost is not None else None,
        "formula": (
            f"R$ {Decimal(str(v_un)):.2f} ÷ {factor_f:g}"
            if v_un is not None and factor_f is not None and factor_f > 1
            else None
        ),
        "packagingStatus": pack.get("packaging_status"),
    }
    return {
        "xmlStatus": kind,
        "xmlStatusLabel": STATUS_PT.get(kind, kind),
        "candidates": len(hits),
        "chosen": {
            "documentId": chosen["documentId"],
            "nNF": chosen.get("nNF"),
            "serie": chosen.get("serie"),
            "issuedAt": chosen.get("issuedAt"),
            "issuerName": chosen.get("issuerName"),
            "issuerCnpj": chosen.get("issuerCnpj"),
            "accessKeyMasked": chosen.get("accessKeyMasked"),
            "cEAN": it.get("c_ean"),
            "cEANTrib": it.get("c_ean_trib"),
            "xProd": it.get("x_prod"),
            "cProd": it.get("c_prod"),
            "ncm": it.get("ncm"),
            "cest": it.get("cest"),
            "origem": it.get("origem"),
            "cfop": it.get("cfop"),
            "cstCsosn": it.get("cst_csosn"),
            "uCom": it.get("u_com"),
            "uTrib": it.get("u_trib"),
            "qCom": it.get("q_com"),
            "qTrib": it.get("q_trib"),
            "vUnCom": it.get("v_un_com"),
            "conversionFactor": factor,
            "purchaseUnitCost": unit_cost,
            "packaging": pack,
            "explain": explain,
            "inboundTaxEvidence": {
                "icms": (it.get("normalized_json") or {}).get("icms"),
                "pis": (it.get("normalized_json") or {}).get("pis"),
                "cofins": (it.get("normalized_json") or {}).get("cofins"),
                "note": "evidência de entrada — não aplicar automaticamente na saída",
            },
        },
        "history": [
            {
                "documentId": h["documentId"],
                "nNF": h.get("nNF"),
                "issuedAt": h.get("issuedAt"),
                "vUnCom": h["item"].get("v_un_com"),
            }
            for h in sorted(hits, key=lambda x: str(x.get("issuedAt") or ""), reverse=True)[:5]
        ],
    }


def _catalog_price(p: Any, table: str = "A") -> str | None:
    attr = {"A": "preco_a", "B": "preco_b", "C": "preco_c"}.get((table or "A").upper(), "preco_a")
    val = getattr(p, attr, None) if not isinstance(p, dict) else p.get(attr) or p.get("precoA")
    return str(val) if val is not None else None


def _match_webposto(
    line: dict[str, Any], snap: dict[str, Any] | None, *, price_table: str = "A"
) -> dict[str, Any]:
    if line.get("eanIssues") and any(str(i).startswith("BLOCKED_") for i in line["eanIssues"]):
        return {
            "wpStatus": "INVALID_EAN",
            "wpStatusLabel": STATUS_PT["INVALID_EAN"],
            "matches": [],
        }
    if snap is None:
        return {
            "wpStatus": "NEW_PRODUCT",
            "wpStatusLabel": STATUS_PT["NEW_PRODUCT"],
            "matches": [],
            "snapshotError": True,
        }

    by_ean = snap.get("byEan") or {}
    hits = list(by_ean.get(line.get("ean") or "", []) or [])
    norm_hits = []
    for h in hits:
        if hasattr(h, "produto_codigo"):
            norm_hits.append(
                {
                    "produtoCodigo": h.produto_codigo,
                    "descricao": fix_mojibake(getattr(h, "nome", None)),
                    "eans": getattr(h, "eans", []) or [],
                    "ativo": getattr(h, "ativo", True),
                    "preco": _catalog_price(h, price_table),
                }
            )
        elif isinstance(h, dict):
            norm_hits.append(
                {
                    "produtoCodigo": h.get("produtoCodigo") or h.get("codigo"),
                    "descricao": fix_mojibake(h.get("descricao") or h.get("nome")),
                    "eans": h.get("eans") or [],
                    "ativo": h.get("ativo", True),
                    "preco": str(h.get("preco") or h.get("precoA") or "") or None,
                }
            )

    if not norm_hits and line.get("descricaoNorm"):
        products_raw = snap.get("products") or {}
        products = list(products_raw.values()) if isinstance(products_raw, dict) else list(products_raw)
        name_hits = []
        for p in products:
            nome = getattr(p, "nome", None) if not isinstance(p, dict) else p.get("nome") or p.get("descricao")
            desc = normalize_name(nome or "")
            if desc and desc == line["descricaoNorm"]:
                name_hits.append(p)
        if len(name_hits) > 1:
            return {
                "wpStatus": "POSSIBLE_DUPLICATE",
                "wpStatusLabel": STATUS_PT["POSSIBLE_DUPLICATE"],
                "matches": [],
                "reason": "várias descrições iguais",
            }
        if len(name_hits) == 1:
            p = name_hits[0]
            if hasattr(p, "produto_codigo"):
                norm_hits = [
                    {
                        "produtoCodigo": p.produto_codigo,
                        "descricao": fix_mojibake(p.nome),
                        "ativo": p.ativo,
                        "preco": _catalog_price(p, price_table),
                    }
                ]
            else:
                norm_hits = [
                    {
                        "produtoCodigo": p.get("produtoCodigo"),
                        "descricao": fix_mojibake(p.get("nome") or p.get("descricao")),
                        "ativo": p.get("ativo", True),
                        "preco": str(p.get("precoA") or p.get("preco") or "") or None,
                    }
                ]

    if len(norm_hits) > 1:
        return {
            "wpStatus": "AMBIGUOUS",
            "wpStatusLabel": STATUS_PT["AMBIGUOUS"],
            "matches": norm_hits[:5],
        }
    if len(norm_hits) == 1:
        m = norm_hits[0]
        ativo = m.get("ativo", True)
        st = "ALREADY_REGISTERED" if ativo else "INACTIVE_EXISTING"
        price_diff = None
        try:
            if line.get("precoVenda") and m.get("preco"):
                price_diff = str(Decimal(line["precoVenda"]) - Decimal(str(m["preco"])))
        except Exception:
            price_diff = None
        return {
            "wpStatus": st,
            "wpStatusLabel": STATUS_PT[st],
            "matches": [m],
            "produtoCodigo": m.get("produtoCodigo"),
            "precoAtual": m.get("preco"),
            "precoDiff": price_diff,
        }
    return {"wpStatus": "NEW_PRODUCT", "wpStatusLabel": STATUS_PT["NEW_PRODUCT"], "matches": []}


def _match_fiscal_model(nfe_chosen: dict[str, Any] | None, models: list[dict[str, Any]]) -> dict[str, Any]:
    if not nfe_chosen:
        return {"fiscalStatus": "NO_FISCAL_MODEL", "model": None, "suggestion": None}
    ncm = nfe_chosen.get("ncm")
    cest = nfe_chosen.get("cest")
    approved = [
        m
        for m in models
        if m.get("status") == "APPROVED"
        and str(m.get("uf", "")).upper() == DEFAULT_UF
        and "PRESUMIDO" in str(m.get("tax_regime", "")).upper().replace(" ", "_")
    ]
    if not approved:
        approved = [
            m
            for m in models
            if m.get("status") == "APPROVED" and str(m.get("uf", "")).upper() == DEFAULT_UF
        ]
    exact = [m for m in approved if ncm and m.get("ncm") == ncm and (not cest or not m.get("cest") or m.get("cest") == cest)]
    if exact:
        m = exact[0]
        fields = {f.get("field_name"): f.get("value_json") for f in (m.get("fields") or [])}
        missing = [
            k
            for k in (
                "tributoIcms",
                "tributoPisCofins",
                "naturezaReceitaCodigo",
                "cdCfopEntrada",
                "cdCfopSaida",
                "iat",
                "ippt",
            )
            if not fields.get(k)
        ]
        return {
            "fiscalStatus": "TAX_CONFLICT" if missing else "MODEL_OK",
            "model": {
                "id": m.get("id"),
                "name": m.get("name"),
                "category": m.get("category"),
                "ncm": m.get("ncm"),
                "cest": m.get("cest"),
                "group_code": m.get("group_code"),
                "cost_center_code": m.get("cost_center_code") or DEFAULT_COST_CENTER,
                "missingOutboundFields": missing,
            },
            "suggestion": None,
        }

    # sugestão draft / similar — não libera
    similar = [m for m in models if ncm and m.get("ncm") == ncm][:1]
    return {
        "fiscalStatus": "NO_FISCAL_MODEL",
        "fiscalStatusLabel": STATUS_PT["NO_FISCAL_MODEL"],
        "model": None,
        "suggestion": {
            "id": similar[0].get("id"),
            "name": similar[0].get("name"),
            "status": similar[0].get("status"),
            "ncm": similar[0].get("ncm"),
            "note": "sugestão — não confirmado / não APPROVED",
        }
        if similar
        else None,
    }


def _final_status(line: dict[str, Any], wp: dict, xml: dict, fiscal: dict) -> tuple[str, list[str]]:
    missing: list[str] = []
    if line.get("duplicateInSheetRow"):
        return "POSSIBLE_DUPLICATE", ["duplicado na planilha"]
    if not line.get("ean") or "BLOCKED_" in str(line.get("eanIssues")):
        if any(str(i).startswith("BLOCKED_") for i in (line.get("eanIssues") or [])):
            return "INVALID_ROW", line.get("eanIssues") or ["EAN inválido"]
    if not line.get("descricao"):
        missing.append("descricao")
    if not line.get("precoVenda"):
        missing.append("precoVenda")

    wp_st = wp.get("wpStatus")
    if wp_st in {"ALREADY_REGISTERED", "INACTIVE_EXISTING"}:
        return wp_st, []
    if wp_st in {"AMBIGUOUS", "POSSIBLE_DUPLICATE"}:
        return "POSSIBLE_DUPLICATE", ["conflito no webPosto"]
    if wp_st == "INVALID_EAN":
        return "INVALID_ROW", ["EAN inválido"]

    xml_st = xml.get("xmlStatus")
    if xml_st == "XML_AMBIGUOUS":
        return "READY_FOR_REVIEW", ["vários candidatos NF-e"]
    if xml_st == "XML_NOT_FOUND":
        return "NO_XML_EVIDENCE", ["sem evidência NF-e"]

    chosen = xml.get("chosen") or {}
    pack_st = (chosen.get("packaging") or {}).get("packaging_status")
    if pack_st == "PACKAGING_AMBIGUITY":
        return "PACKAGING_AMBIGUITY", ["embalagem ambígua"]

    if chosen.get("purchaseUnitCost") is None:
        return "MISSING_PURCHASE_PRICE", ["sem custo unitário"]

    if fiscal.get("fiscalStatus") == "TAX_CONFLICT":
        return "TAX_CONFLICT", fiscal.get("model", {}).get("missingOutboundFields") or ["campos fiscais"]
    if fiscal.get("fiscalStatus") != "MODEL_OK":
        return "NO_FISCAL_MODEL", ["sem modelo APPROVED"]

    model = fiscal.get("model") or {}
    if not chosen.get("ncm"):
        missing.append("ncm")
    if not model.get("group_code"):
        missing.append("grupoCodigo")
    if missing:
        return "INCOMPLETE_BODY", missing

    # READY_TO_CREATE
    if (
        wp_st == "NEW_PRODUCT"
        and xml_st in {"XML_EXACT_MATCH", "XML_PACKAGING_MATCH"}
        and pack_st in {"UNIT_CONFIRMED", "PACKAGE_CONFIRMED", "CONVERSION_CONFIRMED"}
        and fiscal.get("fiscalStatus") == "MODEL_OK"
        and line.get("precoVenda")
        and line.get("descricao")
        and line.get("ean")
    ):
        return "READY_TO_CREATE", []

    return "READY_FOR_REVIEW", ["revisão manual"]


def _load_webposto_snapshot(company_code: int) -> dict[str, Any] | None:
    try:
        from src.operational.price_update.product_snapshot_service import ProductSnapshotService

        snap = ProductSnapshotService().build(company_code)
        return {
            "byEan": snap.get("byEan") or {},
            "products": snap.get("products") or {},
            "productCount": snap.get("productCount"),
        }
    except Exception as exc:
        return {"byEan": {}, "products": {}, "error": type(exc).__name__}


def analyze_registration_spreadsheet(
    *,
    company_code: int,
    company_cnpj: str,
    spreadsheet: bytes,
    filename: str,
    price_table: str = "A",
    sheet: str | None = None,
    col_ean: str | None = None,
    col_desc: str | None = None,
    col_price: str | None = None,
    actor: str = "ui",
) -> dict[str, Any]:
    parsed = parse_target_spreadsheet(
        spreadsheet,
        filename,
        sheet=sheet,
        col_ean=col_ean,
        col_desc=col_desc,
        col_price=col_price,
    )
    nfe_idx = _load_nfe_index(company_code)
    snap = _load_webposto_snapshot(company_code)
    try:
        from src.operational.fiscal_models.service import list_models

        models = list_models(company_code=company_code)
    except Exception:
        models = []

    rows_out: list[dict[str, Any]] = []
    counts: dict[str, int] = {}
    probe: dict[str, Any] | None = None

    for line in parsed["lines"]:
        wp = _match_webposto(line, snap, price_table=price_table)
        if snap and snap.get("error"):
            wp["snapshotWarning"] = snap["error"]
        xml = _match_nfe(line, nfe_idx)
        fiscal = _match_fiscal_model(xml.get("chosen"), models)
        final, missing = _final_status(line, wp, xml, fiscal)
        counts[final] = counts.get(final, 0) + 1

        chosen = xml.get("chosen") or {}
        detail_xml = {k: v for k, v in xml.items() if k != "chosen"}
        detail_xml["chosen"] = chosen
        row = {
            "row": line["row"],
            "eanPlanilha": line.get("ean"),
            "descricao": line.get("descricao"),
            "precoVenda": line.get("precoVenda"),
            "webpostoStatus": wp.get("wpStatus"),
            "webpostoStatusLabel": wp.get("wpStatusLabel"),
            "produtoCodigo": wp.get("produtoCodigo"),
            "precoAtualWebposto": wp.get("precoAtual"),
            "precoDiff": wp.get("precoDiff"),
            "xmlStatus": xml.get("xmlStatus"),
            "xmlStatusLabel": xml.get("xmlStatusLabel"),
            "eanNfe": chosen.get("cEANTrib") or chosen.get("cEAN"),
            "eanCaixa": chosen.get("cEAN") if chosen.get("cEAN") != chosen.get("cEANTrib") else None,
            "fornecedor": chosen.get("issuerName"),
            "ultimaNfe": f"{chosen.get('nNF') or '—'}/{chosen.get('serie') or '—'}",
            "issuedAt": chosen.get("issuedAt"),
            "uCom": chosen.get("uCom"),
            "uTrib": chosen.get("uTrib"),
            "conversao": chosen.get("conversionFactor"),
            "precoCompraUnitario": chosen.get("purchaseUnitCost"),
            "packagingExplain": chosen.get("explain"),
            "ncm": chosen.get("ncm"),
            "cest": chosen.get("cest"),
            "modeloFiscal": (fiscal.get("model") or {}).get("name") or (fiscal.get("suggestion") or {}).get("name"),
            "modeloFiscalId": (fiscal.get("model") or {}).get("id") or (fiscal.get("suggestion") or {}).get("id"),
            "modeloFiscalStatus": (fiscal.get("model") and "APPROVED")
            or ((fiscal.get("suggestion") or {}).get("status")),
            "centroCusto": DEFAULT_COST_CENTER,
            "tabelaPreco": price_table,
            "camposAusentes": missing,
            "status": final,
            "statusLabel": STATUS_PT.get(final, final),
            "actions": _actions_for(final),
            "detail": {
                "webposto": wp,
                "xml": detail_xml,
                "fiscal": fiscal,
            },
        }
        rows_out.append(row)

        if (
            probe is None
            and final == "READY_TO_CREATE"
            and xml.get("xmlStatus") == "XML_EXACT_MATCH"
            and not _looks_fuel(line.get("descricao") or "")
        ):
            probe = {
                "row": line["row"],
                "ean": line.get("ean"),
                "descricao": line.get("descricao"),
                "precoVenda": line.get("precoVenda"),
                "precoCompraUnitario": chosen.get("purchaseUnitCost"),
                "ncm": chosen.get("ncm"),
                "cest": chosen.get("cest"),
                "modeloFiscalId": (fiscal.get("model") or {}).get("id"),
                "note": "candidato a probe — POST bloqueado nesta etapa",
            }

    report = {
        "success": True,
        "empresa": {
            "empresaCodigo": company_code,
            "nome": "CONVENIENCIA 24 HORAS",
            "cnpj": company_cnpj,
            "uf": DEFAULT_UF,
            "regime": "LUCRO PRESUMIDO",
            "centroCustoCodigo": DEFAULT_COST_CENTER,
            "centroCustoNome": "CONVENIENCIA",
            "tabelaPreco": price_table,
        },
        "planilha": {
            "nome": parsed["filename"],
            "aba": parsed["sheet"],
            "tamanho": parsed["size"],
            "sha256": parsed["sha256"],
            "headers": parsed["headers"],
            "mapping": parsed["mapping"],
            "linhas": parsed["lineCount"],
        },
        "nfeReutilizadas": nfe_idx["documentCount"],
        "itensFiscaisPesquisados": nfe_idx["itemCount"],
        "novoUploadZip": False,
        "counts": {
            "ALREADY_REGISTERED": counts.get("ALREADY_REGISTERED", 0) + counts.get("INACTIVE_EXISTING", 0),
            "NEW_PRODUCTS": sum(1 for r in rows_out if r["webpostoStatus"] == "NEW_PRODUCT"),
            "READY_TO_CREATE": counts.get("READY_TO_CREATE", 0),
            "READY_FOR_REVIEW": counts.get("READY_FOR_REVIEW", 0),
            "POSSIBLE_DUPLICATES": counts.get("POSSIBLE_DUPLICATE", 0),
            "NO_XML_EVIDENCE": counts.get("NO_XML_EVIDENCE", 0),
            "NO_FISCAL_MODEL": counts.get("NO_FISCAL_MODEL", 0),
            "PACKAGING_AMBIGUITY": counts.get("PACKAGING_AMBIGUITY", 0),
            "TAX_CONFLICT": counts.get("TAX_CONFLICT", 0),
            "BLOCKED": counts.get("BLOCKED", 0) + counts.get("INVALID_ROW", 0) + counts.get("INCOMPLETE_BODY", 0) + counts.get("MISSING_PURCHASE_PRICE", 0),
            "LINHAS": parsed["lineCount"],
        },
        "statusBreakdown": counts,
        "rows": rows_out,
        "probeCandidate": probe,
        "statusLabels": STATUS_PT,
        "message": (
            f"Serão utilizadas as {nfe_idx['documentCount']} NF-e e os {nfe_idx['itemCount']} "
            "itens fiscais já importados. Nenhum produto foi cadastrado."
        ),
        "webpostoWrites": 0,
        "sefazQueries": 0,
        "analyzedAt": _now(),
        "actor": actor,
    }

    # persist report (caminho absoluto no store DF-e)
    import json

    out_dir = store.store_root() / "registration_reports"
    out_dir.mkdir(parents=True, exist_ok=True)
    report_id = f"reg_{sha256_bytes(spreadsheet)[:12]}_{datetime.now(timezone.utc).strftime('%Y%m%d%H%M%S')}"
    path = out_dir / f"{report_id}.json"

    clean = fix_tree(report)
    clean["reportId"] = report_id
    clean["analysisStatus"] = "PRODUCT_SPREADSHEET_ANALYZED"
    path.write_text(json.dumps(clean, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    clean["reportPath"] = str(path.resolve())

    store.append_audit(
        company_code=company_code,
        action="REGISTRATION_ANALYSIS",
        actor=actor,
        target_type="report",
        target_id=report_id,
        detail={"linhas": parsed["lineCount"], "sha256": parsed["sha256"], "ready": counts.get("READY_TO_CREATE", 0)},
    )
    return clean


def _actions_for(status: str) -> list[str]:
    if status in {"ALREADY_REGISTERED", "INACTIVE_EXISTING"}:
        return ["VER PRODUTO"]
    if status == "POSSIBLE_DUPLICATE":
        return ["COMPARAR"]
    if status == "PACKAGING_AMBIGUITY":
        return ["REVISAR EMBALAGEM"]
    if status == "NO_FISCAL_MODEL":
        return ["PROCURAR MODELO", "CRIAR MODELO FISCAL", "ENVIAR PARA REVISÃO"]
    if status == "READY_TO_CREATE":
        return ["VER BODY DO CADASTRO"]
    return ["ENVIAR PARA REVISÃO"]


def _looks_fuel(desc: str) -> bool:
    u = desc.upper()
    return any(x in u for x in ("GASOLINA", "ETANOL", "DIESEL", "S10", "S500", "GNV", "COMBUST"))


def dfe_inventory_summary(company_code: int) -> dict[str, Any]:
    idx = _load_nfe_index(company_code)
    return {
        "empresaCodigo": company_code,
        "nfeImportadas": idx["documentCount"],
        "itensDisponiveis": idx["itemCount"],
        "message": (
            f"Serão utilizadas as {idx['documentCount']} NF-e e os {idx['itemCount']} "
            "itens fiscais já importados."
        ),
        "novoUploadZip": False,
        "analysisStatus": "WAITING_FOR_USER_SPREADSHEET",
        "webpostoWrites": 0,
        "sefazQueries": 0,
    }
