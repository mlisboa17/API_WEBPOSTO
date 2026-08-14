"""Geração de XLSX e JSON de relatórios — FASE 10."""

from __future__ import annotations

import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

from .schemas import ProductAnalysis

logger = logging.getLogger(__name__)


class ReportExporter:
    """Gera relatórios XLSX e JSON."""

    def __init__(self, output_dir: Path | str = "reports/product_registration"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_preflight_xlsx(
        self,
        all_analyses: list[ProductAnalysis],
        file_hash: str,
        execution_id: str,
    ) -> Path:
        """
        Exporta XLSX pré-voo com:
        - SUMMARY: contagens
        - READY_TO_CREATE: produtos prontos
        - BLOCKED: bloqueados
        - REVIEW_REQUIRED: pendentes revisão
        """
        
        filepath = self.output_dir / f"cadastro_preflight_{execution_id}.xlsx"
        
        # Agrupar por status
        by_status: dict[str, list[ProductAnalysis]] = {}
        for analysis in all_analyses:
            if analysis.status not in by_status:
                by_status[analysis.status] = []
            by_status[analysis.status].append(analysis)

        # Contar por risco
        by_risk: dict[str, int] = {
            "BAIXO_RISCO": len([a for a in all_analyses if a.risk_level == "BAIXO_RISCO"]),
            "MÉDIO_RISCO": len([a for a in all_analyses if a.risk_level == "MÉDIO_RISCO"]),
            "ALTO_RISCO": len([a for a in all_analyses if a.risk_level == "ALTO_RISCO"]),
        }

        wb = openpyxl.Workbook()
        
        # Aba SUMMARY
        ws_summary = wb.active
        ws_summary.title = "SUMMARY"
        ws_summary.append(["Métrica", "Valor"])
        ws_summary.append(["TOTAL", len(all_analyses)])
        ws_summary.append(["READY_TO_CREATE", len(by_status.get("READY_TO_CREATE", []))])
        ws_summary.append(["BLOCKED", len(by_status.get("BLOCKED", []))])
        ws_summary.append(["REVIEW_REQUIRED", len(by_status.get("REVIEW_REQUIRED", []))])
        ws_summary.append(["BAIXO_RISCO", by_risk["BAIXO_RISCO"]])
        ws_summary.append(["MÉDIO_RISCO", by_risk["MÉDIO_RISCO"]])
        ws_summary.append(["ALTO_RISCO", by_risk["ALTO_RISCO"]])
        ws_summary.append(["File Hash", file_hash])
        ws_summary.append(["Execution ID", execution_id])

        # Aba READY_TO_CREATE
        self._add_analysis_sheet(
            wb,
            "READY_TO_CREATE",
            by_status.get("READY_TO_CREATE", []),
        )

        # Aba BLOCKED
        self._add_analysis_sheet(
            wb,
            "BLOCKED",
            by_status.get("BLOCKED", []),
        )

        # Aba REVIEW_REQUIRED
        self._add_analysis_sheet(
            wb,
            "REVIEW_REQUIRED",
            by_status.get("REVIEW_REQUIRED", []),
        )

        wb.save(filepath)
        logger.info(f"[EXPORT] Pré-voo XLSX: {filepath}")
        return filepath

    def _add_analysis_sheet(
        self,
        wb: openpyxl.Workbook,
        sheet_name: str,
        analyses: list[ProductAnalysis],
    ) -> None:
        """Adiciona aba de análises."""
        ws = wb.create_sheet(sheet_name)
        ws.append([
            "EAN",
            "Descrição",
            "Preço Venda",
            "Status",
            "Gate",
            "Confiança",
            "Risco",
            "Grupo",
            "Centro",
            "NCM",
            "CEST",
            "Problemas",
        ])

        for analysis in analyses:
            ws.append([
                analysis.ean,
                analysis.descricao,
                analysis.preco_venda,
                analysis.status,
                analysis.gate,
                analysis.confidence,
                analysis.risk_level,
                analysis.grupo_api_codigo,
                analysis.centro_api_codigo,
                analysis.ncm,
                analysis.cest,
                " | ".join(analysis.validation_issues),
            ])

    def export_execution_xlsx(
        self,
        created_verified: list[dict[str, Any]],
        created_not_verified: list[dict[str, Any]],
        rejected: list[dict[str, Any]],
        result_unknown: list[dict[str, Any]],
        execution_id: str,
    ) -> Path:
        """
        Exporta XLSX de execução com:
        - CREATED_AND_VERIFIED
        - CREATED_BUT_NOT_VERIFIED
        - REJECTED
        - RESULT_UNKNOWN
        """
        
        filepath = self.output_dir / f"cadastro_execucao_{execution_id}.xlsx"

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        self._add_execution_sheet(wb, "CREATED_AND_VERIFIED", created_verified)
        self._add_execution_sheet(wb, "CREATED_BUT_NOT_VERIFIED", created_not_verified)
        self._add_execution_sheet(wb, "REJECTED", rejected)
        self._add_execution_sheet(wb, "RESULT_UNKNOWN", result_unknown)

        wb.save(filepath)
        logger.info(f"[EXPORT] Execução XLSX: {filepath}")
        return filepath

    def _add_execution_sheet(
        self,
        wb: openpyxl.Workbook,
        sheet_name: str,
        records: list[dict[str, Any]],
    ) -> None:
        """Adiciona aba de execução."""
        ws = wb.create_sheet(sheet_name)
        
        headers = [
            "EAN",
            "Descrição",
            "HTTP Status",
            "RET",
            "MEN",
            "CodProduto",
            "Body Hash",
            "Criado Em",
        ]
        ws.append(headers)

        for record in records:
            ws.append([
                record.get("ean"),
                record.get("descricao"),
                record.get("http_status"),
                record.get("ret"),
                record.get("men", "")[:100],
                record.get("cod_produto"),
                record.get("body_hash", "")[:8],
                record.get("created_at"),
            ])

    def export_relatorio_json(
        self,
        summary: dict[str, Any],
        execution_id: str,
    ) -> Path:
        """Exporta relatório JSON final."""
        
        filepath = self.output_dir / f"cadastro_relatorio_{execution_id}.json"

        report = {
            "title": "PRODUCT BULK REGISTRATION REPORT",
            "timestamp": datetime.now().isoformat(),
            "execution_id": execution_id,
            **summary,
        }

        filepath.write_text(json.dumps(report, indent=2, ensure_ascii=False))
        logger.info(f"[EXPORT] Relatório JSON: {filepath}")
        return filepath

    def export_auditoria_jsonl(
        self,
        records: list[dict[str, Any]],
        execution_id: str,
    ) -> Path:
        """Exporta auditoria JSONL (uma linha por produto processado)."""
        
        filepath = self.output_dir / f"cadastro_auditoria_{execution_id}.jsonl"

        with open(filepath, "w", encoding="utf-8") as f:
            for record in records:
                f.write(json.dumps(record, ensure_ascii=False) + "\n")

        logger.info(f"[EXPORT] Auditoria JSONL: {filepath}")
        return filepath

    def export_successful_templates_json(
        self,
        templates: dict[str, Any],
        execution_id: str,
    ) -> Path:
        """Exporta templates reutilizáveis."""
        
        filepath = self.output_dir / f"successful_templates_{execution_id}.json"

        filepath.write_text(json.dumps(templates, indent=2, ensure_ascii=False))
        logger.info(f"[EXPORT] Templates JSON: {filepath}")
        return filepath
